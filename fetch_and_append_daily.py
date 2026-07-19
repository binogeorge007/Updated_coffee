"""
Best-effort daily fetch + append for Coffee_Price_Log.xlsx.

Fetches:
  - Karnataka local Arabica/Robusta prices from kirehalli.com
  - ICE Arabica futures from tradingeconomics.com (as a proxy/benchmark)
  - AUD/INR and USD/INR from xe.com

Appends 4 rows (Arabica Parchment/Cherry, Robusta Parchment/Cherry) to the
"Price Log" sheet of Coffee_Price_Log.xlsx (must already exist - run
build_log_local.py first).

CAVEATS - read before relying on this:
  - This scrapes public web pages with regex/BeautifulSoup. These sites do not
    provide a stable API and can change their HTML at any time, which will
    silently break the parsing below. Check the printed output each run.
  - Some sites may rate-limit or block repeated automated requests from a
    single IP over time. There's no guarantee this keeps working long-term.
  - Nothing here fabricates a number: if a fetch fails, that field is left
    as "N/A" and printed as a warning rather than guessed.
  - This is NOT connected to any scheduler. It only runs when you execute it.
    To run it automatically every morning, add it to cron (Mac/Linux) or
    Task Scheduler (Windows) - see the bottom of this file for examples.

Requires: pip install requests beautifulsoup4 openpyxl
"""
import re
import sys
import datetime
import requests
from bs4 import BeautifulSoup
import openpyxl

WORKBOOK_PATH = "Coffee_Price_Log.xlsx"
HEADERS = {"User-Agent": "Mozilla/5.0 (compatible; CoffeePriceBot/1.0)"}

FREIGHT_ROW, CUSTOMS_ROW, WAREHOUSE_ROW, DELIVERY_ROW, MARGIN_ROW = 5, 6, 7, 8, 9


def get(url, timeout=15):
    try:
        resp = requests.get(url, headers=HEADERS, timeout=timeout)
        resp.raise_for_status()
        return resp.text
    except Exception as e:
        print(f"  [warn] fetch failed for {url}: {e}")
        return None


def fetch_kirehalli():
    """Find the latest Kirehalli 'Coffee Prices (Karnataka)' post and parse it."""
    html = get("https://kirehalli.com/coffee-prices-daily/")
    if not html:
        return None
    soup = BeautifulSoup(html, "html.parser")
    link = None
    for a in soup.find_all("a", href=True):
        if "coffee-prices-karnataka" in a["href"]:
            link = a["href"]
            break
    if not link:
        print("  [warn] could not find a Kirehalli post link")
        return None

    post_html = get(link)
    if not post_html:
        return None
    text = BeautifulSoup(post_html, "html.parser").get_text("\n")

    result = {"source_url": link}

    def find_range(label):
        m = re.search(rf"{label}.*?Rs\s*([\d,]+)\s*[–-]\s*(?:Rs\s*)?([\d,]+)", text, re.IGNORECASE | re.DOTALL)
        if not m:
            return None, None
        return int(m.group(1).replace(",", "")), int(m.group(2).replace(",", ""))

    result["arabica_parchment"] = find_range("Arabica Parchment")
    result["arabica_cherry"] = find_range("Arabica Cherry")
    result["robusta_parchment"] = find_range("Robusta Parchment")
    result["robusta_cherry"] = find_range("Robusta Cherry")

    m = re.search(r"arabica coffee.*?([\d]+\.[\d]+)\s*(?:US\s*)?cents?/lb", text, re.IGNORECASE)
    result["ice_arabica_cents_lb"] = float(m.group(1)) if m else None

    m = re.search(r"robusta coffee.*?US\$?\s*([\d,]+)\s*/?\s*tonne", text, re.IGNORECASE)
    result["ice_robusta_usd_tonne"] = float(m.group(1).replace(",", "")) if m else None

    return result


def fetch_ice_arabica_benchmark():
    html = get("https://tradingeconomics.com/commodity/coffee")
    if not html:
        return None
    m = re.search(r"rose to ([\d.]+) USd/Lbs|fell to ([\d.]+) USd/Lbs|to ([\d.]+) USd/Lbs", html)
    if not m:
        return None
    val = next(g for g in m.groups() if g)
    return float(val)


def fetch_rate(from_ccy, to_ccy):
    html = get(f"https://www.xe.com/currencyconverter/convert/?Amount=1&From={from_ccy}&To={to_ccy}")
    if not html:
        return None
    m = re.search(rf"1\.00 {from_ccy} = ([\d.]+) {to_ccy}", html)
    if not m:
        return None
    return float(m.group(1))


def main():
    print("Fetching Karnataka local prices (Kirehalli)...")
    kirehalli = fetch_kirehalli() or {}

    print("Fetching ICE Arabica benchmark (Trading Economics)...")
    ice_arabica_benchmark = fetch_ice_arabica_benchmark()

    print("Fetching AUD/INR and USD/INR (xe.com)...")
    aud_inr = fetch_rate("AUD", "INR")
    usd_inr = fetch_rate("USD", "INR")

    ice_arabica = kirehalli.get("ice_arabica_cents_lb") or ice_arabica_benchmark
    ice_robusta = kirehalli.get("ice_robusta_usd_tonne")

    print("\n--- Fetched values ---")
    print("Arabica Parchment:", kirehalli.get("arabica_parchment"))
    print("Arabica Cherry:", kirehalli.get("arabica_cherry"))
    print("Robusta Parchment:", kirehalli.get("robusta_parchment"))
    print("Robusta Cherry:", kirehalli.get("robusta_cherry"))
    print("ICE Arabica (c/lb):", ice_arabica)
    print("ICE Robusta ($/tonne):", ice_robusta)
    print("AUD/INR:", aud_inr, " USD/INR:", usd_inr)

    try:
        wb = openpyxl.load_workbook(WORKBOOK_PATH)
    except FileNotFoundError:
        print(f"\n[error] {WORKBOOK_PATH} not found. Run build_log_local.py first.")
        sys.exit(1)
    ws = wb["Price Log"]
    next_row = ws.max_row + 1
    today = datetime.date.today()

    rows_to_add = [
        ("Arabica", "Parchment (AP)", kirehalli.get("arabica_parchment"), ice_arabica, "US cents/lb (ICE Arabica)"),
        ("Arabica", "Cherry (AC)", kirehalli.get("arabica_cherry"), ice_arabica, "US cents/lb (ICE Arabica)"),
        ("Robusta", "Parchment (RP)", kirehalli.get("robusta_parchment"), ice_robusta, "USD/tonne (ICE Robusta)"),
        ("Robusta", "Cherry (RC)", kirehalli.get("robusta_cherry"), ice_robusta, "USD/tonne (ICE Robusta)"),
    ]

    r = next_row
    for variety, grade, price_range, ice_price, ice_unit in rows_to_add:
        low, high = (price_range if price_range else (None, None))
        ws.cell(row=r, column=1, value=today)
        ws.cell(row=r, column=1).number_format = "yyyy-mm-dd"
        ws.cell(row=r, column=2, value=variety)
        ws.cell(row=r, column=3, value=grade)
        ws.cell(row=r, column=4, value=low if low is not None else "N/A")
        ws.cell(row=r, column=5, value=high if high is not None else "N/A")
        ws.cell(row=r, column=6, value="")
        ws.cell(row=r, column=7, value="Kirehalli (auto-fetched)")
        ws.cell(row=r, column=8, value=ice_price if ice_price is not None else "N/A")
        ws.cell(row=r, column=9, value=ice_unit)
        ws.cell(row=r, column=10, value=aud_inr if aud_inr is not None else "N/A")
        ws.cell(row=r, column=11, value=usd_inr if usd_inr is not None else "N/A")
        if low is not None and aud_inr:
            # Landed cost off the MILLED-equivalent price (what's actually
            # exported/cleared through customs), not the raw parchment/cherry
            # price - see build_log_local.py comments for why.
            outturn_expr = (
                f'IF(B{r}="Arabica",IF(ISNUMBER(SEARCH("Parchment",C{r})),0.8,0.48),'
                f'IF(ISNUMBER(SEARCH("Parchment",C{r})),0.82,0.5))'
            )
            milled_base = f'((AVERAGE(D{r}:E{r})/50/J{r})/({outturn_expr}))'
            formula = (
                f"={milled_base}*(1+Assumptions!$B${CUSTOMS_ROW})"
                f"+Assumptions!$B${FREIGHT_ROW}+Assumptions!$B${WAREHOUSE_ROW}+Assumptions!$B${DELIVERY_ROW}"
                f"+{milled_base}*Assumptions!$B${MARGIN_ROW}"
            )
            ws.cell(row=r, column=12, value=formula)
        ws.cell(row=r, column=16, value="Auto-fetched" if low is not None else "Fetch failed - fill in manually")

        # Mid price / weeks tracked / tracked high-low / range status - only if we have a real price today
        if low is not None:
            ws.cell(row=r, column=17, value=f"=AVERAGE(D{r}:E{r})")
            ws.cell(row=r, column=18, value=f"=(A{r}-_xlfn.MINIFS($A$2:A{r},$B$2:B{r},B{r},$C$2:C{r},C{r}))/7")
            ws.cell(row=r, column=18).number_format = "0.0"
            ws.cell(row=r, column=19, value=f"=_xlfn.MAXIFS($Q$2:Q{r},$B$2:B{r},B{r},$C$2:C{r},C{r})")
            ws.cell(row=r, column=20, value=f"=_xlfn.MINIFS($Q$2:Q{r},$B$2:B{r},B{r},$C$2:C{r},C{r})")
            ws.cell(row=r, column=21, value=(
                f'=IF(R{r}<1,"Day 1 - building history",'
                f'IF(Q{r}>=S{r},TEXT(ROUND(R{r},0),"0")&"-wk high (since tracking began)",'
                f'IF(Q{r}<=T{r},TEXT(ROUND(R{r},0),"0")&"-wk low (since tracking began)",'
                f'TEXT(ROUND(R{r},0),"0")&"-wk range, mid-range")))'
            ))
        if low is not None and aud_inr and ice_price is not None and usd_inr:
            ws.cell(row=r, column=22, value=f"=Q{r}/50/K{r}")
            ws.cell(row=r, column=23, value=f'=IF(B{r}="Arabica",H{r}/100/0.45359237,H{r}/1000)')

        # X: 3/6-Month Flag - green (low, safer to buy) / red (high, riskier),
        # computed from this sheet's own accumulated history. Only meaningful
        # once 90+/180+ days of real logged history exist for that variety+grade.
        if low is not None:
            hist = f"(A{r}-_xlfn.MINIFS($A$2:A{r},$B$2:B{r},B{r},$C$2:C{r},C{r}))"
            six_low = f'_xlfn.MINIFS($Q$2:Q{r},$B$2:B{r},B{r},$C$2:C{r},C{r},$A$2:A{r},">="&(A{r}-180))'
            six_high = f'_xlfn.MAXIFS($Q$2:Q{r},$B$2:B{r},B{r},$C$2:C{r},C{r},$A$2:A{r},">="&(A{r}-180))'
            three_low = f'_xlfn.MINIFS($Q$2:Q{r},$B$2:B{r},B{r},$C$2:C{r},C{r},$A$2:A{r},">="&(A{r}-90))'
            three_high = f'_xlfn.MAXIFS($Q$2:Q{r},$B$2:B{r},B{r},$C$2:C{r},C{r},$A$2:A{r},">="&(A{r}-90))'
            ws.cell(row=r, column=24, value=(
                f'=IF(AND({hist}>=180,Q{r}<={six_low}),"6-MONTH LOW - historically cheap, safer to buy",'
                f'IF(AND({hist}>=180,Q{r}>={six_high}),"6-MONTH HIGH - historically expensive, riskier to buy",'
                f'IF(AND({hist}>=90,Q{r}<={three_low}),"3-MONTH LOW - historically cheap, safer to buy",'
                f'IF(AND({hist}>=90,Q{r}>={three_high}),"3-MONTH HIGH - historically expensive, riskier to buy",'
                f'IF({hist}>=90,"within 3/6-month range","N/A - insufficient history yet (need 90+ days)")))))'
            ))

            # Y/Z: milled/clean-bean-equivalent USD/kg + outturn % used (same
            # industry rule-of-thumb ratios as build_log_local.py - see its
            # comments for the caveat on precision).
            outturn_expr = (
                f'IF(B{r}="Arabica",IF(ISNUMBER(SEARCH("Parchment",C{r})),0.8,0.48),'
                f'IF(ISNUMBER(SEARCH("Parchment",C{r})),0.82,0.5))'
            )
            ws.cell(row=r, column=25, value=f"=V{r}/{outturn_expr}")
            ws.cell(row=r, column=26, value=(
                f'=IF(B{r}="Arabica",IF(ISNUMBER(SEARCH("Parchment",C{r})),"80%","48%"),'
                f'IF(ISNUMBER(SEARCH("Parchment",C{r})),"82%","50%"))'
            ))
            ws.cell(row=r, column=27, value=(
                f'=IF(B{r}="Arabica",IF(ISNUMBER(SEARCH("Parchment",C{r})),'
                f'"Arabica Plantation (estimated milled)","Arabica Cherry (estimated milled)"),'
                f'IF(ISNUMBER(SEARCH("Parchment",C{r})),"Robusta Parchment (estimated milled)",'
                f'"Robusta Cherry (estimated milled)"))'
            ))
        r += 1

    # Re-apply conditional formatting over the full column so newly appended
    # rows are covered (green = LOW flag, red = HIGH flag).
    from openpyxl.styles import PatternFill
    from openpyxl.formatting.rule import FormulaRule
    ws.conditional_formatting._cf_rules.clear()  # avoid stacking duplicate rules on every run
    flag_range = f"X2:X{ws.max_row}"
    green_fill = PatternFill("solid", fgColor="C6EFCE")
    red_fill = PatternFill("solid", fgColor="FFC7CE")
    ws.conditional_formatting.add(flag_range, FormulaRule(formula=['ISNUMBER(SEARCH("LOW",X2))'], fill=green_fill))
    ws.conditional_formatting.add(flag_range, FormulaRule(formula=['ISNUMBER(SEARCH("HIGH",X2))'], fill=red_fill))

    wb.save(WORKBOOK_PATH)
    print(f"\nAppended {r - next_row} rows to {WORKBOOK_PATH}.")
    print("Open it in Excel to let formulas recalculate (Excel does this automatically on open).")


if __name__ == "__main__":
    main()

# ----------------------------------------------------------------------------
# To run this automatically every morning at 9am:
#
# Mac/Linux (cron): run `crontab -e` and add a line like:
#   0 9 * * * cd /path/to/folder && /usr/bin/python3 fetch_and_append_daily.py >> log.txt 2>&1
#
# Windows (Task Scheduler): create a Basic Task, trigger "Daily" at 9:00 AM,
#   action "Start a program", program = python.exe, arguments =
#   "fetch_and_append_daily.py", start-in = the folder containing this file.
# ----------------------------------------------------------------------------
